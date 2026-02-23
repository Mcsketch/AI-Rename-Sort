"""File content extraction for various file types."""
import base64
import io
import os
from pathlib import Path

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tiff", ".tif"}
PDF_EXTENSIONS = {".pdf"}
VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".wmv", ".flv", ".webm", ".m4v"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".log", ".py", ".js", ".ts", ".html", ".xml",
    ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".sh",
    ".bat", ".css", ".scss", ".rs", ".go", ".java", ".c", ".cpp", ".h",
}
DOCUMENT_EXTENSIONS = {".docx", ".xlsx", ".pptx", ".doc", ".xls", ".ppt", ".odt", ".ods", ".odp"}

MAX_TEXT_LENGTH = 4000
MAX_IMAGE_DIMENSION = 1024
MAX_VIDEO_THUMBNAIL_SIZE = 512
MAX_PDF_PAGES = 10
MAX_EXCEL_SHEETS = 3
MAX_EXCEL_ROWS = 20


class FileProcessor:
    """Extracts content from files for AI analysis."""

    def get_file_type(self, filepath):
        """Return a string category for the given file path."""
        ext = Path(filepath).suffix.lower()
        if ext in IMAGE_EXTENSIONS:
            return "image"
        if ext in PDF_EXTENSIONS:
            return "pdf"
        if ext in VIDEO_EXTENSIONS:
            return "video"
        if ext in TEXT_EXTENSIONS:
            return "text"
        if ext in DOCUMENT_EXTENSIONS:
            return "document"
        return "unknown"

    def extract_content(self, filepath):
        """Return (content, file_type) for the given file.

        ``content`` is either a base64 data URL (for images) or a plain string.
        """
        file_type = self.get_file_type(filepath)
        if file_type == "image":
            return self._extract_image(filepath)
        if file_type == "pdf":
            return self._extract_pdf(filepath)
        if file_type == "video":
            return self._extract_video(filepath)
        if file_type == "text":
            return self._extract_text(filepath)
        if file_type == "document":
            return self._extract_document(filepath)
        # Unknown: return basic file info
        size = os.path.getsize(filepath)
        return f"Unknown file: {Path(filepath).name}\nSize: {size / 1024:.1f} KB", "unknown"

    # ------------------------------------------------------------------
    # Private helpers
    # ------------------------------------------------------------------

    def _extract_image(self, filepath):
        """Encode image as a base64 JPEG data URL (resized for API limits)."""
        try:
            from PIL import Image

            with Image.open(filepath) as img:
                if img.mode not in ("RGB", "RGBA", "L"):
                    img = img.convert("RGB")
                elif img.mode == "RGBA":
                    img = img.convert("RGB")
                img.thumbnail((MAX_IMAGE_DIMENSION, MAX_IMAGE_DIMENSION), Image.Resampling.LANCZOS)
                buf = io.BytesIO()
                img.save(buf, format="JPEG")
                encoded = base64.b64encode(buf.getvalue()).decode("utf-8")
            return f"data:image/jpeg;base64,{encoded}", "image"
        except ImportError:
            # Pillow not available – fall back to raw bytes
            ext = Path(filepath).suffix.lower().lstrip(".")
            mime_map = {
                "jpg": "image/jpeg", "jpeg": "image/jpeg",
                "png": "image/png", "gif": "image/gif",
                "bmp": "image/bmp", "webp": "image/webp",
                "tiff": "image/tiff", "tif": "image/tiff",
            }
            mime = mime_map.get(ext, "image/jpeg")
            with open(filepath, "rb") as f:
                encoded = base64.b64encode(f.read()).decode("utf-8")
            return f"data:{mime};base64,{encoded}", "image"
        except Exception as exc:
            return f"Image file (could not read: {exc})", "image"

    def _extract_pdf(self, filepath):
        """Extract text from a PDF using pdfplumber."""
        try:
            import pdfplumber

            parts = []
            with pdfplumber.open(filepath) as pdf:
                for i, page in enumerate(pdf.pages[:MAX_PDF_PAGES]):
                    text = page.extract_text()
                    if text:
                        parts.append(f"[Page {i + 1}]\n{text}")
            content = "\n\n".join(parts)
            return (content[:MAX_TEXT_LENGTH] if content else "Empty PDF"), "pdf"
        except ImportError:
            return "PDF file (pdfplumber not installed)", "pdf"
        except Exception as exc:
            return f"PDF file (could not read: {exc})", "pdf"

    def _extract_video(self, filepath):
        """Extract a representative frame from a video file, if possible."""
        try:
            import cv2
            from PIL import Image as PILImage

            cap = cv2.VideoCapture(str(filepath))
            total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            if total > 0:
                cap.set(cv2.CAP_PROP_POS_FRAMES, max(0, total // 10))
            ret, frame = cap.read()
            cap.release()
            if ret:
                frame_rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
                img = PILImage.fromarray(frame_rgb)
                img.thumbnail((MAX_VIDEO_THUMBNAIL_SIZE, MAX_VIDEO_THUMBNAIL_SIZE))
                buf = io.BytesIO()
                img.save(buf, format="JPEG")
                encoded = base64.b64encode(buf.getvalue()).decode("utf-8")
                return f"data:image/jpeg;base64,{encoded}", "image"
        except Exception:
            pass

        # Fallback: describe the file
        size = os.path.getsize(filepath)
        info_lines = [f"Video file: {Path(filepath).name}", f"Size: {size / (1024 * 1024):.1f} MB"]
        duration = self._get_video_duration(filepath)
        if duration:
            info_lines.append(duration)
        return "\n".join(info_lines), "video"

    def _get_video_duration(self, filepath):
        import subprocess

        try:
            result = subprocess.run(
                [
                    "ffprobe", "-v", "error",
                    "-show_entries", "format=duration",
                    "-of", "default=noprint_wrappers=1:nokey=1",
                    str(filepath),
                ],
                capture_output=True,
                text=True,
                timeout=10,
            )
            if result.returncode == 0:
                secs = float(result.stdout.strip())
                return f"Duration: {int(secs // 60)}m {int(secs % 60)}s"
        except Exception:
            pass
        return ""

    def _extract_text(self, filepath):
        """Read a plain-text file."""
        try:
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                content = f.read(MAX_TEXT_LENGTH)
            return content, "text"
        except Exception as exc:
            return f"Text file (could not read: {exc})", "text"

    def _extract_document(self, filepath):
        """Extract text from Office-format documents."""
        ext = Path(filepath).suffix.lower()

        if ext == ".docx":
            try:
                import docx

                doc = docx.Document(filepath)
                content = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
                return content[:MAX_TEXT_LENGTH], "document"
            except ImportError:
                pass
            except Exception as exc:
                return f"Word document (could not read: {exc})", "document"

        if ext in (".xlsx", ".xls"):
            try:
                import openpyxl

                wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
                parts = []
                for sheet in list(wb.worksheets)[:MAX_EXCEL_SHEETS]:
                    rows = []
                    for row in sheet.iter_rows(max_row=MAX_EXCEL_ROWS, values_only=True):
                        row_str = ", ".join(str(c) for c in row if c is not None)
                        if row_str:
                            rows.append(row_str)
                    if rows:
                        parts.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
                return "\n\n".join(parts)[:MAX_TEXT_LENGTH], "document"
            except ImportError:
                pass
            except Exception as exc:
                return f"Excel document (could not read: {exc})", "document"

        size = os.path.getsize(filepath)
        return (
            f"Document file: {Path(filepath).name}\nSize: {size / 1024:.1f} KB\nType: {ext}",
            "document",
        )
